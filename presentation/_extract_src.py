import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
print("SLIDES:", len(p.slides))
for i, sl in enumerate(p.slides, 1):
    print(f"\n===== SLIDE {i} =====")
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip()
            print("[TXT]", t.replace("\n", " || "))
        if sh.has_table:
            print("[TABLE]")
            for row in sh.table.rows:
                print("   |", " | ".join(c.text.strip() for c in row.cells))
    if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
        print("[NOTES]", sl.notes_slide.notes_text_frame.text.strip()[:200])
