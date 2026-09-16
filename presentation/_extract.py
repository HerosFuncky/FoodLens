from pptx import Presentation
from pptx.util import Emu
p = Presentation("FoodLens_presentation.pptx")
print("Slides:", len(p.slides))
for i, sl in enumerate(p.slides, 1):
    texts = []
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n"," | "))
    print(f"\n=== Slide {i} ===")
    for t in texts:
        print("  -", t[:95])
